"""
Verificador "Agora não" - levantamento paralelo ao Sender (não envia nada).

Pesquisa "agora não" no WhatsApp Web, desce a lista de resultados até o fim
e junta todos os contatos cuja mensagem é exatamente a resposta do botão
"Agora não". Depois cruza com a planilha e, se você confirmar, move esses
contatos para a aba "Não enviar" (com backup da planilha antes).

O robô só LÊ a tela: não abre conversas, não clica em contatos, não envia nada.
Uso:  python verificar_agora_nao.py      (com o Chrome do Sender FECHADO)
"""

import re
import shutil
import sys
import time
from datetime import datetime
from pathlib import Path

import openpyxl
from selenium.common.exceptions import WebDriverException
from selenium.webdriver.common.by import By
from selenium.webdriver.common.keys import Keys

import sender  # reaproveita configuração, login e validação de números do Sender

TERMO_BUSCA = "agora não"
RESPOSTA_EXATA = "agora não"          # só conta quem mandou exatamente isso
ABA_NAO_ENVIAR = "Não enviar"
MOTIVO = "Respondeu \"Agora não\""
ARQUIVO_RESULTADO = str(Path(sender.ARQUIVO_REGISTRO).with_name("Agora não.md"))
PAUSA_SCROLL = 1.5                    # segundos para o WhatsApp carregar mais resultados
PAUSA_NO_FIM = 4                      # no fim da lista, espera mais (a busca carrega em lotes)
TENTATIVAS_NO_FIM = 5                 # só termina após N tentativas no fim sem nada novo

LINHA = sender.LINHA

# Lê os resultados da busca direto da página. Cada resultado vira [linhas de texto].
# Procura pelas linhas da lista (role=row/listitem) e, se o WhatsApp mudar o layout,
# cai no plano B: sobe a partir de cada trecho de mensagem até achar o "cartão".
# Acha o painel que REALMENTE rola: sobe a partir de um resultado da busca até o
# primeiro elemento com barra de rolagem. (Na busca, o WhatsApp não usa o mesmo
# painel da lista de conversas — rolar o painel errado parecia "fim da lista".)
JS_PAINEL = r"""
window.__painelBusca = (function () {
  const rolavel = e => e && e.scrollHeight > e.clientHeight + 20 &&
                       /(auto|scroll|overlay)/.test(getComputedStyle(e).overflowY);
  const lateral = document.querySelector('#side') ? document.querySelector('#side').parentElement : document.body;
  const linha = lateral.querySelector('[role="row"], [role="listitem"], [data-testid="last-msg-status"]');
  for (let e = linha; e && e !== document.body; e = e.parentElement) if (rolavel(e)) return e;
  let melhor = null;
  for (const e of lateral.querySelectorAll('*'))
    if (rolavel(e) && (!melhor || e.clientHeight > melhor.clientHeight)) melhor = e;
  return melhor;
})();
const p = window.__painelBusca;
return p ? {ok: true, id: p.id || p.className.toString().slice(0, 40),
            altura: p.clientHeight, total: p.scrollHeight} : {ok: false};
"""

JS_LER_RESULTADOS = r"""
const painel = window.__painelBusca || document.querySelector('#pane-side') || document.querySelector('#side');
if (!painel) return {erro: 'painel'};
let linhas = [...painel.querySelectorAll('[role="row"], [role="listitem"]')];
if (!linhas.length) {
  linhas = [...painel.querySelectorAll('[data-testid="last-msg-status"], span[title]')]
    .map(s => { let e = s; for (let i = 0; i < 8 && e; i++) {
                  if (e.innerText && e.innerText.split('\n').length >= 3) return e; e = e.parentElement; }
                return null; })
    .filter(Boolean);
}
const limpa = t => t.replace(/[\u202a-\u202e\u200e\u200f]/g, '').trim();
return {itens: [...new Set(linhas)].map(l => l.innerText.split('\n').map(limpa).filter(Boolean))};
"""

JS_ROLAR = r"""
const painel = window.__painelBusca || document.querySelector('#pane-side') || document.querySelector('#side');
const antes = painel.scrollTop;
painel.scrollTop = antes + painel.clientHeight * 0.7;
painel.dispatchEvent(new Event('scroll', {bubbles: true}));
return {antes: antes, depois: painel.scrollTop, total: painel.scrollHeight,
        fim: painel.scrollTop + painel.clientHeight >= painel.scrollHeight - 5};
"""


def normalizar(texto):
    return re.sub(r"\s+", " ", texto).strip().lower()


def pesquisar(driver):
    """Digita o termo na barra de pesquisa da lista de conversas."""
    candidatos = driver.find_elements(By.CSS_SELECTOR,
                                      "#side div[contenteditable='true'], #side input[type='text'], "
                                      "div[contenteditable='true'][data-tab='3']")
    if not candidatos:
        raise RuntimeError("Não encontrei a barra de pesquisa do WhatsApp.")
    barra = candidatos[0]
    barra.click()
    barra.send_keys(Keys.CONTROL, "a")
    barra.send_keys(Keys.DELETE)
    barra.send_keys(TERMO_BUSCA)
    sender.esperar(5, "Carregando a pesquisa")


def coletar(driver):
    """Desce a lista de resultados até o fim. Devolve {titulo: trecho}."""
    achados = {}
    try:
        _coletar(driver, achados)
    except sender.ParadaEmergencia:
        print("\n  ESC - parando a coleta e usando o que já foi encontrado.")
    return achados


def ler_tela(driver, achados):
    """Lê os resultados visíveis e acrescenta os novos. Devolve quantos eram novos."""
    dados = driver.execute_script(JS_LER_RESULTADOS)
    if dados.get("erro"):
        raise RuntimeError("Não encontrei a lista de resultados na tela.")
    antes = len(achados)
    for linhas in dados["itens"]:
        if len(linhas) < 2:
            continue
        titulo = linhas[0]
        if sender.bloqueado(numero_do_titulo(titulo)):          # lista "Nunca contatar"
            continue
        if any(normalizar(l) == RESPOSTA_EXATA for l in linhas[1:]):
            achados.setdefault(titulo, linhas[-1])
    return len(achados) - antes


def _coletar(driver, achados):
    painel = driver.execute_script(JS_PAINEL)
    if not painel.get("ok"):
        raise RuntimeError("Não encontrei o painel de resultados que rola.")
    print(f"  Painel da busca: {painel['altura']}px visíveis de {painel['total']}px")

    rolagens = 0
    tentativas_no_fim = 0
    while tentativas_no_fim < TENTATIVAS_NO_FIM:
        ler_tela(driver, achados)
        rolagem = driver.execute_script(JS_ROLAR)
        rolagens += 1
        aguardar(PAUSA_SCROLL)
        novos = ler_tela(driver, achados)
        if rolagem["fim"] and rolagem["depois"] == rolagem["antes"] and not novos:
            # no fim da lista: dá tempo para o WhatsApp carregar o próximo lote
            tentativas_no_fim += 1
            aguardar(PAUSA_NO_FIM)
        else:
            tentativas_no_fim = 0
        print(f"\r  Rolagem {rolagens:>4} | encontrados: {len(achados):>4} | "
              f"posição {rolagem['depois']:>6}/{rolagem['total']:<6}", end="", flush=True)
    ler_tela(driver, achados)
    print()


def aguardar(segundos):
    """Espera silenciosa que ainda responde ao ESC."""
    fim = time.time() + segundos
    while time.time() < fim:
        if sender.msvcrt.kbhit() and sender.msvcrt.getwch() == "\x1b":
            raise sender.ParadaEmergencia
        time.sleep(0.1)


def numero_do_titulo(titulo):
    """'+55 15 99123-4567' -> '5515991234567'. Nome de contato salvo -> None."""
    if not re.fullmatch(r"[+\d\s().-]{8,}", titulo):
        return None
    return sender.formatar_numero(titulo) or re.sub(r"\D", "", titulo)


def cruzar_com_planilha(numeros):
    import pandas as pd
    df = sender.ler_planilha()
    na_planilha = []
    for i, linha in df.iterrows():
        numero = sender.formatar_numero(linha[sender.COLUNA_NUMERO])
        if numero and numero in numeros:
            na_planilha.append({"linha": sender.linha_excel(i), "nome": sender.nome_da_planilha(linha[sender.COLUNA_NOME]),
                                "numero": numero})
    return na_planilha


def ler_registro_anterior():
    """{contato: data em que foi visto pela 1ª vez} do 'Agora não.md' que já existe."""
    anteriores = {}
    arquivo = Path(ARQUIVO_RESULTADO)
    if not arquivo.exists():
        return anteriores
    for linha in arquivo.read_text(encoding="utf-8").splitlines():
        celulas = [c.strip() for c in linha.strip().strip("|").split("|")]
        if len(celulas) >= 4 and re.fullmatch(r"\d{2}/\d{2}/\d{4}", celulas[-1]):
            if sender.bloqueado(numero_do_titulo(celulas[0])):   # lista "Nunca contatar"
                continue
            anteriores[celulas[0]] = celulas[-1]
    return anteriores


def salvar_resultado(achados, na_planilha):
    """Registro ACUMULADO 'Agora não.md' no Markdown: quem já foi registrado
    continua lá; quem aparecer pela primeira vez entra com a data de hoje."""
    anteriores = ler_registro_anterior()
    hoje = f"{datetime.now():%d/%m/%Y}"
    todos = {t: anteriores.get(t, hoje) for t in anteriores | dict.fromkeys(achados)}
    novos = [t for t in achados if t not in anteriores]

    com_numero = {t: numero_do_titulo(t) for t in todos}
    linhas_planilha = {c["numero"]: c for c in na_planilha}
    texto = [f"# 🙅 Responderam \"Agora não\"\n",
             "> Registro acumulado de todos que responderam exatamente \"Agora não\" "
             "(botão de resposta rápida de uma campanha), levantado pesquisando no WhatsApp Web.\n"
             f"> Última verificação: {datetime.now():%d/%m/%Y às %H:%M}.\n",
             f"- **Total registrado:** {len(todos)}",
             f"- **Novos nesta verificação:** {len(novos)}",
             f"- **Estão na planilha agora:** {len(na_planilha)}",
             f"- **Contatos salvos com nome (conferir manualmente):** "
             f"{sum(1 for n in com_numero.values() if n is None)}\n",
             "| Contato no WhatsApp | Número | Na planilha? | Visto em |",
             "|---|---|---|---|"]
    ordem = sorted(todos, key=lambda t: (com_numero[t] is None, t))
    for titulo in ordem:
        numero = com_numero[titulo]
        c = linhas_planilha.get(numero)
        na = f"✅ linha {c['linha']} ({c['nome']})" if c else ("❓ conferir" if numero is None else "—")
        texto.append(f"| {titulo.replace('|', '/')} | {numero or '—'} | {na} | {todos[titulo]} |")
    Path(ARQUIVO_RESULTADO).write_text("\n".join(texto) + "\n", encoding="utf-8")
    return len(todos), len(novos)


def mover_para_nao_enviar(na_planilha):
    """Backup + tira as linhas da aba principal e guarda na aba 'Não enviar'."""
    planilha = Path(sender.PLANILHA)
    backup = planilha.with_name(f"{planilha.stem} - backup {datetime.now():%Y-%m-%d %H%M}{planilha.suffix}")
    shutil.copy2(planilha, backup)
    print(f"  Backup salvo: {backup.name}")

    livro = openpyxl.load_workbook(planilha)
    principal = livro.worksheets[sender.ABA]
    if ABA_NAO_ENVIAR in livro.sheetnames:
        destino = livro[ABA_NAO_ENVIAR]
    else:
        destino = livro.create_sheet(ABA_NAO_ENVIAR)
        cabecalho = [c.value for c in principal[sender.LINHA_CABECALHO + 1]]
        destino.append(cabecalho + ["Motivo", "Movido em"])

    cabecalho = [c.value for c in principal[sender.LINHA_CABECALHO + 1]]
    col_numero = cabecalho.index(sender.COLUNA_NUMERO)
    movidos = 0
    for c in sorted(na_planilha, key=lambda c: c["linha"], reverse=True):   # de baixo pra cima
        valores = [cel.value for cel in principal[c["linha"]]]
        # confere que a linha é mesmo desse contato antes de tirar da planilha
        if sender.formatar_numero(str(valores[col_numero] or "")) != c["numero"]:
            print(f"  ⚠️  Linha {c['linha']} não bate com {c['numero']} - não mexi nela.")
            continue
        destino.append(valores + [MOTIVO, datetime.now().strftime("%d/%m/%Y %H:%M")])
        principal.delete_rows(c["linha"])
        movidos += 1
    livro.save(planilha)
    print(f"  ✅ {movidos} contatos movidos para a aba \"{ABA_NAO_ENVIAR}\".")


def main():
    print(LINHA)
    print("  Verificador \"Agora não\" - só lê o WhatsApp, não envia nada")
    print(LINHA)
    print("  Antes de começar, FECHE o Chrome do Sender (ele usa o mesmo login).")
    input("  Aperte Enter quando estiver pronto...")

    driver = None
    try:
        driver = sender.abrir_whatsapp()
        pesquisar(driver)
        print("\n  Descendo a lista de resultados (ESC para parar e usar o que já achou)...")
        achados = coletar(driver)
    except (WebDriverException, RuntimeError) as erro:
        print(f"\n  ❌ {str(erro).splitlines()[0][:150]}")
        return
    except sender.ParadaEmergencia:
        return
    finally:
        if driver is not None:
            try:
                driver.quit()
            except WebDriverException:
                pass

    # cruza com TODOS os já registrados, não só os desta busca
    titulos = set(achados) | set(ler_registro_anterior())
    numeros = {n for n in map(numero_do_titulo, titulos) if n}
    na_planilha = cruzar_com_planilha(numeros)
    total, novos = salvar_resultado(achados, na_planilha)

    print(f"\n{LINHA}")
    print(f"  Encontrados nesta busca:      {len(achados)}")
    print(f"  Registro acumulado:           {total}  ({novos} novos)")
    print(f"    - com número:               {len(numeros)}")
    print(f"    - contato salvo com nome:   {total - len(numeros)}  (conferir manualmente)")
    print(f"  Estão na planilha:            {len(na_planilha)}")
    for c in na_planilha:
        print(f"    linha {c['linha']:>3} - {c['nome']} - {c['numero']}")
    print(f"  Lista completa no Markdown: Sender/{Path(ARQUIVO_RESULTADO).name}")
    print(LINHA)

    if not na_planilha:
        return
    resposta = sender.perguntar(
        f"\n  Mover esses {len(na_planilha)} contatos para a aba \"{ABA_NAO_ENVIAR}\"? "
        "(faço backup antes)\n  [1] Sim, mover   [0] Não: ", ["1", "0"])
    if resposta == "1":
        try:
            mover_para_nao_enviar(na_planilha)
        except PermissionError:
            print("  ❌ A planilha está aberta no Excel. Feche o Excel e rode de novo.")


if __name__ == "__main__":
    try:
        main()
    except KeyboardInterrupt:
        print("\n  Cancelado.")
        sys.exit(1)
